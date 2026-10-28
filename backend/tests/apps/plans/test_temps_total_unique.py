"""
#759 — « aucune ventilation » : le temps de travail se réduit à un seul
« Temps de travail total » par année, sans distinction agent / partenaire.

Quand l'action n'est ni ventilée ni déclinée par poste, les lots sans cible
(un financé, un non financé…) fusionnent en une seule ligne financée, au
prévisionnel comme au réalisé. Les autres modes gardent leurs lots.
"""

import io
from decimal import Decimal

from openpyxl import load_workbook

import pytest
from rest_framework.test import APIClient

from apps.plans.models_operations import (
    OperationAnnee, OperationAnneeRH, RealisationOperationAnnee, RealisationOperationAnneeRH,
)
from apps.plans.rh_total import est_temps_total_unique, fusionner_temps_total
from apps.plans.services_export_fiche_action import build_fiche_action_workbook
from tests.factories.enjeux import (
    EnjeuFactory, IndicateurFactory, MetriqueFactory, NiveauExigenceFactory,
    ObjectifLongTermeFactory, OperationFactory,
)
from tests.factories.plans import PlanGestionFactory
from tests.factories.users import OrganismeFactory, SuperAdminFactory


def _client():
    client = APIClient()
    client.force_authenticate(SuperAdminFactory())
    return client


@pytest.fixture
def action(db):
    """Action 2028 d'un plan brouillon, sans déclinaison par type de coût."""
    plan = PlanGestionFactory(annee_debut=2027, annee_fin=2030, statut='draft')
    enjeu = EnjeuFactory(id_pg=plan)
    ne = NiveauExigenceFactory(id_olt=ObjectifLongTermeFactory(id_enjeu=enjeu))
    met = MetriqueFactory(id_indicateur=IndicateurFactory(id_ne=ne))

    def make(mode):
        op = OperationFactory(
            metriques=[met], ventilation_mode=mode, declinaison_par_type_cout=False,
            annee_min=2028, annee_max=2028,
        )
        oa = OperationAnnee.objects.create(id_operation=op, annee=2028, periodicite=True)
        return op, oa
    return make


def _deux_lots():
    return [
        {'jours': '25', 'categorie_depense': 'fonctionnement'},
        {'jours': '6', 'categorie_depense': 'benevolat_partenariat'},
    ]


@pytest.mark.django_db
class TestProgrammation:
    def test_aucune_ventilation_une_seule_ligne(self, action):
        op, _oa = action('none')
        resp = _client().patch(f'/api/plans/operations/{op.pk}/', {
            'operation_annees': [{'annee': 2028, 'periodicite': True, 'organismes': [],
                                  'rh_lignes': _deux_lots()}],
        }, format='json')
        assert resp.status_code == 200, resp.content
        lignes = OperationAnneeRH.objects.filter(id_operation_annee__id_operation=op)
        assert lignes.count() == 1
        ligne = lignes.get()
        assert ligne.jours == Decimal('31')
        assert ligne.finance is True
        assert ligne.categorie_depense == 'fonctionnement'

        # Fiche action (Excel / Word) : plus de ligne bénévoles / partenaires.
        wb = load_workbook(io.BytesIO(build_fiche_action_workbook(op.get_plan_de_gestion())))
        textes = [c.value for ws in wb.worksheets
                  for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert 'Jours bénévoles / partenaires' not in textes
        assert 'Travail prévisionnel (jours)' in textes

    def test_par_type_de_budget_garde_les_lots(self, action):
        op, _oa = action('by_type')
        resp = _client().patch(f'/api/plans/operations/{op.pk}/', {
            'operation_annees': [{'annee': 2028, 'periodicite': True, 'organismes': [],
                                  'rh_lignes': _deux_lots()}],
        }, format='json')
        assert resp.status_code == 200, resp.content
        assert OperationAnneeRH.objects.filter(id_operation_annee__id_operation=op).count() == 2


@pytest.mark.django_db
class TestSuivi:
    def test_realise_fusionne_et_reste_rattache_au_prevu(self, action):
        op, oa = action('none')
        prevu = OperationAnneeRH.objects.create(id_operation_annee=oa, jours=31)
        resp = _client().post('/api/plans/realisations/upsert/', {
            'id_operation_annee': oa.id_operation_annee,
            'rh_lignes': [
                {'jours': '20', 'finance': True},
                {'jours': '4', 'finance': False},
            ],
        }, format='json')
        assert resp.status_code == 200, resp.content
        lignes = RealisationOperationAnneeRH.objects.filter(
            id_realisation_operation_annee__id_operation_annee=oa)
        assert lignes.count() == 1
        ligne = lignes.get()
        assert ligne.jours == Decimal('24')
        assert ligne.finance is True
        assert ligne.id_operation_annee_rh_id == prevu.pk


@pytest.mark.django_db
class TestDonneesExistantes:
    """Ce que fait la migration 0137 : fusion des saisies antérieures."""

    def test_fusion_prevu_et_realise(self, action):
        op, oa = action('none')
        agent = OperationAnneeRH.objects.create(id_operation_annee=oa, jours=25)
        partenaire = OperationAnneeRH.objects.create(
            id_operation_annee=oa, jours=6, categorie_depense='benevolat_partenariat')
        real = RealisationOperationAnnee.objects.create(id_operation_annee=oa)
        RealisationOperationAnneeRH.objects.create(
            id_realisation_operation_annee=real, id_operation_annee_rh=agent, jours=20)
        RealisationOperationAnneeRH.objects.create(
            id_realisation_operation_annee=real, id_operation_annee_rh=partenaire, jours=4,
            categorie_depense='benevolat_partenariat')

        fusionner_temps_total(op)

        prevus = list(OperationAnneeRH.objects.filter(id_operation_annee=oa))
        assert [(p.jours, p.finance) for p in prevus] == [(Decimal('31'), True)]
        reels = list(real.rh_lignes.all())
        assert [(r.jours, r.finance, r.id_operation_annee_rh_id) for r in reels] == [
            (Decimal('24'), True, prevus[0].pk)]

    def test_lignes_avec_cible_intactes(self, action):
        """Une ligne visant un poste ou un organisme n'est pas un temps total."""
        op, oa = action('none')
        OperationAnneeRH.objects.create(id_operation_annee=oa, jours=10)
        OperationAnneeRH.objects.create(id_operation_annee=oa, jours=2)
        cible = OperationAnneeRH.objects.create(
            id_operation_annee=oa, jours=3, finance=False, id_organisme=OrganismeFactory())
        fusionner_temps_total(op)
        lignes = OperationAnneeRH.objects.filter(id_operation_annee=oa)
        assert lignes.count() == 2
        assert lignes.get(pk=cible.pk).jours == Decimal('3')
        assert lignes.exclude(pk=cible.pk).get().jours == Decimal('12')

    def test_declinaison_par_poste_non_concernee(self, action):
        """Temps décliné par poste (case de la fiche action) : pas fusionné."""
        op, oa = action('none')
        op.declinaison_par_type_cout = True
        op.declinaison_par_poste = True
        op.save()
        OperationAnneeRH.objects.create(id_operation_annee=oa, jours=25)
        OperationAnneeRH.objects.create(id_operation_annee=oa, jours=6, finance=False)
        fusionner_temps_total(op)
        assert OperationAnneeRH.objects.filter(id_operation_annee=oa).count() == 2

    def test_couts_detailles_sans_poste_temps_total(self, action):
        """#773 — les coûts détaillés ne déclinent plus le temps par poste : en
        « aucune ventilation », il reste un seul temps de travail total."""
        op, oa = action('none')
        op.declinaison_par_type_cout = True
        op.save()
        assert est_temps_total_unique(op)
