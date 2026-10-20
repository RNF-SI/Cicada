"""
#726 — Export du suivi des actions : les pictogrammes de réalisation dans les
cases d'année, comme à l'écran, plutôt que leur libellé.

Les pictogrammes se lisent d'un coup d'œil, le texte non. Choix validé : le
pictogramme seul dans la case, et une légende des six statuts en tête du
classeur.
"""
import io

from openpyxl import load_workbook

from apps.plans.services_export_grille import STATUTS_ACTION, build_grille_workbook

STATUTS = ['planned', 'planned-realized', 'planned-partial',
           'planned-not-realized', 'realized-unplanned', 'partial-unplanned']


def _classeur(lignes, legende=None):
    payload = {
        'titre': 'Suivi des actions — Plan test',
        'onglet': 'Réalisation',
        'meta': [['Onglet', 'Réalisation']],
        'entetes': ['Code', 'Action', '2025', '2026'],
        'lignes': lignes,
    }
    if legende is not None:
        payload['legende'] = legende
    return load_workbook(io.BytesIO(build_grille_workbook(payload))).active


def _ancres(ws):
    """(ligne, colonne) 1-based de chaque image du classeur."""
    return {(img.anchor._from.row + 1, img.anchor._from.col + 1) for img in ws._images}


def _ligne_entete(ws):
    return next(r[0].row for r in ws.iter_rows() if r[0].value == 'Code')


def test_les_six_statuts_ont_un_pictogramme():
    assert set(STATUTS_ACTION) == set(STATUTS)


def test_case_d_annee_porte_le_pictogramme_et_pas_le_texte():
    ws = _classeur([{'cellules': [
        'CS1', 'Action A',
        {'t': 'Action prévue et réalisée', 'a': 'planned-realized'},
        {'t': 'Action prévue', 'a': 'planned'},
    ]}])
    ligne = _ligne_entete(ws) + 1
    assert {(ligne, 3), (ligne, 4)} <= _ancres(ws)
    assert ws.cell(ligne, 3).value is None
    assert ws.cell(ligne, 4).value is None
    assert ws.cell(ligne, 2).value == 'Action A'


def test_case_sans_statut_reste_vide_sans_image():
    ws = _classeur([{'cellules': ['CS1', 'Action A', '', {'t': 'Action prévue', 'a': 'planned'}]}])
    ligne = _ligne_entete(ws) + 1
    assert (ligne, 3) not in _ancres(ws)
    assert (ligne, 4) in _ancres(ws)


def test_statut_inconnu_retombe_sur_le_texte():
    ws = _classeur([{'cellules': ['CS1', 'Action A', {'t': 'Autre', 'a': 'inconnu'}, '']}])
    ligne = _ligne_entete(ws) + 1
    assert ws.cell(ligne, 3).value == 'Autre'
    assert not _ancres(ws)


def test_legende_des_six_pictogrammes_avant_le_tableau():
    legende = [{'a': s, 't': f'Libellé {s}'} for s in STATUTS]
    ws = _classeur([{'cellules': ['CS1', 'Action A', '', '']}], legende=legende)
    entete = _ligne_entete(ws)
    lignes_legende = [r for r in range(1, entete) if str(ws.cell(r, 2).value or '').startswith('Libellé ')]
    assert len(lignes_legende) == 6
    assert {(r, 1) for r in lignes_legende} <= _ancres(ws)
    assert ws.cell(lignes_legende[0], 2).value == 'Libellé planned'
