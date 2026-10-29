"""
#727 — export « fiche de suivi » d'une action : prévu / réalisé d'UNE année.

Une fiche (onglet) par action et par année. Depuis les exports du plan, on
choisit l'année : seules les actions dont le suivi est saisi cette année-là sont
exportées. Depuis la saisie du suivi, la fiche de l'action pour l'année ouverte.
"""

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.plans.models_operations import (
    OperationAnnee, OperationAnneeRH, Poste, RealisationOperationAnnee,
    RealisationOperationAnneeRH,
)
from apps.plans.services_export_fiche_suivi import build_fiche_suivi_workbook
from tests.factories.enjeux import (
    EnjeuFactory, IndicateurFactory, MesureFactory, MetriqueFactory, NiveauExigenceFactory,
    NomenclatureNiveauRealisationFactory, NomenclatureTypeIndicateurFactory,
    NomenclatureTypeMetriqueFactory,
    ObjectifLongTermeFactory, ObjectifOperationnelFactory, OperationFactory,
    ResultatAttenduFactory,
)
from tests.factories.plans import PlanGestionFactory
from tests.factories.users import OrganismeFactory, RoleFactory, SuperAdminFactory


@pytest.fixture
def plan(db):
    """2 actions ; seule « Fauche tardive » a un suivi saisi en 2024."""
    p = PlanGestionFactory(annee_debut=2024, annee_fin=2025, statut='valide',
                           nom="Marais de test")
    enjeu = EnjeuFactory(id_pg=p, libelle="Zones humides")
    olt = ObjectifLongTermeFactory(id_enjeu=enjeu, libelle="Maintenir les prairies humides")
    etat = NomenclatureTypeIndicateurFactory(cd_nomenclature='ETAT', mnemonique='ETAT', label='État')
    reponse = NomenclatureTypeIndicateurFactory(
        cd_nomenclature='REPONSE', mnemonique='REPONSE', label='Réponse')
    texte = NomenclatureTypeMetriqueFactory(cd_nomenclature='TEXTE', mnemonique='TEXTE',
                                            label='Texte')
    met_etat = MetriqueFactory(
        id_indicateur=IndicateurFactory(id_ne=NiveauExigenceFactory(id_olt=olt),
                                        type_indicateur=etat,
                                        nom_indicateur="État des prairies"),
        nom_metrique="Cortège floristique", type_metrique=texte, unite='',
        score_1_label="Dégradé", score_2_label="Appauvri", score_3_label="Moyen",
        score_4_label="Bon état", score_5_label="Optimal")
    # #727 — une mesure dans l'année de la fiche, une autre l'année suivante.
    MesureFactory(id_metrique=met_etat, valeur="Bon état", date_mesure=date(2024, 6, 1))
    MesureFactory(id_metrique=met_etat, valeur="Optimal", date_mesure=date(2025, 6, 1))
    oo = ObjectifOperationnelFactory(id_enjeu=enjeu, libelle="Restaurer la fauche tardive")
    met_rep = MetriqueFactory(
        id_indicateur=IndicateurFactory(id_ne=None, nom_indicateur="Surface fauchée",
                                        id_resultat_attendu=ResultatAttenduFactory(id_oo=oo),
                                        type_indicateur=reponse),
        nom_metrique="Hectares fauchés", etat_reference="12")
    MesureFactory(id_metrique=met_rep, valeur="9", date_mesure=date(2024, 12, 31))

    op = OperationFactory(
        metriques=[met_etat, met_rep], libelle="Fauche tardive",
        operateurs="Garde technicien", financeurs="Agence de l'eau",
        ventilation_mode='by_type_poste', declinaison_par_poste=True,
        declinaison_par_type_cout=True, cout_salarial_auto=True,
        annee_min=2024, annee_max=2025,
    )
    poste = Poste.objects.create(id_pg=p, id_organisme=OrganismeFactory(), nombre=1,
                                 cout_jour=300, nom_local="Garde technicien")
    oa = OperationAnnee.objects.create(id_operation=op, annee=2024, periodicite=True,
                                       cout_prestataire=Decimal('500'))
    OperationAnneeRH.objects.create(id_operation_annee=oa, id_poste=poste, jours=10,
                                    categorie_depense='fonctionnement')
    niveau = NomenclatureNiveauRealisationFactory(mnemonique='PARTIEL', label='Partiellement réalisée')
    real = RealisationOperationAnnee.objects.create(
        id_operation_annee=oa, id_niveau_realisation=niveau,
        cout_prestataire_realise=Decimal('400'), operateurs_realises="Garde + stagiaire",
        commentaires="Fauche décalée en août.")
    RealisationOperationAnneeRH.objects.create(
        id_realisation_operation_annee=real, id_poste=poste, jours=8,
        categorie_depense='fonctionnement')

    sans_suivi = OperationFactory(metriques=[met_etat], libelle="Pose de clôtures",
                                  annee_min=2024, annee_max=2024)
    OperationAnnee.objects.create(id_operation=sans_suivi, annee=2024, periodicite=True)
    p._op, p._sans_suivi = op, sans_suivi
    return p


def _texts(data):
    wb = load_workbook(io.BytesIO(data))
    return wb, [[c.value for c in row] for ws in wb.worksheets for row in ws.iter_rows()]


def _row(rows, label):
    found = [r for r in rows if label in r]
    assert found, f"ligne « {label} » absente"
    return found[0]


@pytest.mark.django_db
class TestClasseur:
    def test_une_fiche_par_action_suivie_dans_l_annee(self, plan):
        wb, rows = _texts(build_fiche_suivi_workbook(plan, 2024))
        assert len(wb.sheetnames) == 1
        flat = [c for r in rows for c in r if isinstance(c, str)]
        assert any("FICHE DE SUIVI" in c and "2024" in c for c in flat)
        assert any("Fauche tardive" in c for c in flat)
        assert not any("Pose de clôtures" in c for c in flat)

    def test_annee_sans_suivi(self, plan):
        wb, rows = _texts(build_fiche_suivi_workbook(plan, 2025))
        flat = [c for r in rows for c in r if isinstance(c, str)]
        assert any("Aucun suivi saisi en 2025" in c for c in flat)

    def test_contenu_prevu_realise(self, plan):
        _wb, rows = _texts(build_fiche_suivi_workbook(plan, 2024))
        assert "Partiellement réalisée" in _row(rows, "Niveau de réalisation")
        operateurs = _row(rows, "Opérateurs")
        assert "Garde technicien" in operateurs and "Garde + stagiaire" in operateurs
        # #727 — prévu en D:F (index 3), réalisé en G:I (index 6) : la fiche
        # s'élargit à 9 colonnes pour la grille des scores (6 paliers).
        garde = next(r for r in rows if r[0] == "Garde technicien")
        assert "10" in garde[3] and "8" in garde[6]
        salarial = _row(rows, "Coût salarial — fonctionnement")
        assert salarial[3] == "3 000 €" and salarial[6] == "2 400 €"
        total = _row(rows, "Total budget")
        assert total[3] == "3 500 €" and total[6] == "2 800 €"
        indic = _row(rows, "Hectares fauchés (Surface fauchée)")
        # Valeur cible / valeur saisie, suivies de l'unité de la métrique.
        assert indic[3].startswith("12") and indic[6].startswith("9")
        assert "Fauche décalée en août." in _row(rows, "Commentaires")
        flat = [c for r in rows for c in r if isinstance(c, str)]
        assert any("Maintenir les prairies humides" in c for c in flat)
        # #766 — l'OO ne vient que de l'indicateur de réponse : il n'est plus
        # repris dans la ligne OLT, qui suit la branche de l'indicateur d'état.
        assert not any("Restaurer la fauche tardive" in c for c in flat)

    def test_rappel_de_l_indicateur_d_etat(self, plan):
        """#727 (retour 23/10) — l'indicateur d'état est rappelé en tête, et la
        ligne d'objectif est nommée selon la branche (OLT)."""
        _wb, rows = _texts(build_fiche_suivi_workbook(plan, 2024))
        assert "État des prairies" in _row(rows, "Indicateur d'état")
        assert "Maintenir les prairies humides" in _row(rows, "OLT")

    def test_mois_de_realisation_retire(self, plan):
        _wb, rows = _texts(build_fiche_suivi_workbook(plan, 2024))
        assert not any("Mois de réalisation" in r for r in rows)

    def test_resultat_de_l_indicateur_d_etat(self, plan):
        """#727 — grille des scores : la mesure de l'année de la fiche, rangée
        dans la colonne de son score ; pas celle d'une autre année."""
        _wb, rows = _texts(build_fiche_suivi_workbook(plan, 2024))
        entete = _row(rows, "Indéterminé")
        assert entete[3:9] == ["Indéterminé", "Très mauvais", "Mauvais", "Moyen", "Bon",
                               "Très bon"]
        ligne = _row(rows, "Cortège floristique (État des prairies)")
        assert ligne[7] == "2024 : Bon état"          # colonne « Bon »
        assert not any(isinstance(c, str) and "Optimal" in c for c in ligne)

    def test_resultat_reprend_la_derniere_mesure_anterieure(self, plan):
        """#727 (retour 28/10) — sans mesure l'année de la fiche, la mesure
        de N-1, sinon N-2, etc., datée de son année ; jamais une postérieure."""
        _wb, rows = _texts(build_fiche_suivi_workbook(plan, 2027, operation_ids=[plan._op.pk]))
        ligne = _row(rows, "Cortège floristique (État des prairies)")
        assert "2025 : Optimal" in ligne
        assert not any(isinstance(c, str) and "Aucune mesure" in c for c in ligne)

    def test_resultat_sans_aucune_mesure(self, plan):
        _wb, rows = _texts(build_fiche_suivi_workbook(plan, 2023, operation_ids=[plan._op.pk]))
        ligne = _row(rows, "Cortège floristique (État des prairies)")
        assert "Aucune mesure" in ligne

    def test_fiche_d_une_action_meme_sans_suivi(self, plan):
        wb, rows = _texts(build_fiche_suivi_workbook(plan, 2024, operation_ids=[plan._sans_suivi.pk]))
        assert len(wb.sheetnames) == 1
        assert "Non saisi" in _row(rows, "Niveau de réalisation")


@pytest.mark.django_db
class TestCarte:
    def test_prevu_et_realise_sur_la_meme_carte(self, settings):
        from django.contrib.gis.geos import Polygon
        from apps.plans.services_export_fiche_action import _geom_png
        settings.EXPORT_MAP_TILE_URL = ''  # pas de fond de carte en test
        prevu = Polygon(((5.0, 45.0), (5.01, 45.0), (5.01, 45.01), (5.0, 45.0)), srid=4326)
        realise = Polygon(((5.02, 45.0), (5.03, 45.0), (5.03, 45.01), (5.02, 45.0)), srid=4326)
        png = _geom_png(prevu, realise=realise)
        assert png and png[:4] == b'\x89PNG'


@pytest.mark.django_db
class TestApi:
    def _client(self, user):
        c = APIClient()
        c.force_authenticate(user)
        return c

    def test_export_du_plan_par_annee(self, plan):
        c = self._client(SuperAdminFactory())
        url = f'/api/plans/plans/{plan.pk}/export-fiches-suivi-xlsx/'
        assert c.get(url).status_code == 400
        r = c.get(url, {'annee': 2024})
        assert r.status_code == 200, r.content
        assert 'fiches-suivi' in r['Content-Disposition'] and '2024' in r['Content-Disposition']

    def test_export_d_une_fiche(self, plan):
        c = self._client(SuperAdminFactory())
        r = c.get(f'/api/plans/operations/{plan._op.pk}/export-fiche-suivi-xlsx/', {'annee': 2024})
        assert r.status_code == 200, r.content

    def test_reserve_aux_referents(self, plan):
        c = self._client(RoleFactory())
        r = c.get(f'/api/plans/plans/{plan.pk}/export-fiches-suivi-xlsx/', {'annee': 2024})
        assert r.status_code in (403, 404)
