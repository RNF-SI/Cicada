"""
Tests HTTP des endpoints d'import / export Excel (#478).

Couvre la couche API (`PlanGestionViewSet`) que les tests service ne touchent
pas :
- export du classeur (GET, tout statut) et bon type MIME ;
- validation (dry-run) via multipart → rapport JSON ;
- import réel dans un brouillon (201 + décomptes) ;
- **verrou brouillon** : import refusé (403) sur un plan validé — la permission
  `CanModifyOnlyDraftPlan` s'applique aux endpoints d'écriture ;
- authentification requise ;
- fichier manquant → 400.
"""

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient

from apps.plans.models import PlanGestion
from apps.plans.services_import import build_arborescence_workbook

from tests.factories.users import SuperAdminFactory
from tests.factories.plans import PlanGestionFactory, PlanGestionValideFactory

# Réutilise les helpers de construction d'un plan source riche.
from tests.apps.plans.test_import_arborescence import (
    _base_nomenclatures,
    _build_source_plan,
)

pytestmark = [pytest.mark.django_db, pytest.mark.integration]

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload_from_plan(plan):
    """Construit un fichier .xlsx d'arborescence à partir d'un plan source."""
    content = build_arborescence_workbook(plan=plan)
    return SimpleUploadedFile("arborescence.xlsx", content, content_type=XLSX_MIME)


def _url(plan, suffix):
    return f"/api/plans/plans/{plan.pk}/{suffix}"


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------


def test_export_arborescence_requires_auth():
    plan = PlanGestionFactory()
    resp = APIClient().get(_url(plan, "export-arborescence-xlsx/"))
    assert resp.status_code in (401, 403)


def test_export_arborescence_ok_on_draft():
    user = SuperAdminFactory()
    _base_nomenclatures()
    plan = _build_source_plan(user)
    client = APIClient()
    client.force_authenticate(user=user)

    resp = client.get(_url(plan, "export-arborescence-xlsx/"))
    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME
    assert b"PK" == resp.content[:2]  # signature d'un .xlsx (zip)


def test_export_arborescence_ok_on_validated_plan():
    """L'export (lecture) fonctionne quel que soit le statut du plan."""
    user = SuperAdminFactory()
    plan = PlanGestionValideFactory(id_utilisateur_ajout=user)
    client = APIClient()
    client.force_authenticate(user=user)

    resp = client.get(_url(plan, "export-arborescence-xlsx/?empty=1"))
    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME


def test_export_actions_ok():
    user = SuperAdminFactory()
    _base_nomenclatures()
    plan = _build_source_plan(user)
    client = APIClient()
    client.force_authenticate(user=user)

    resp = client.get(_url(plan, "export-actions-xlsx/"))
    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME


# ---------------------------------------------------------------------------
# Validation (dry-run)
# ---------------------------------------------------------------------------


def test_validate_endpoint_returns_report():
    user = SuperAdminFactory()
    _base_nomenclatures()
    source = _build_source_plan(user)
    target = PlanGestionFactory(id_utilisateur_ajout=user)

    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(
        _url(target, "import-arborescence/validate/"),
        {"file": _upload_from_plan(source)},
        format="multipart",
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["can_import"] is True
    assert body["summary"]["enjeux"] == source.enjeux.count()
    # Rien n'a été écrit (dry-run).
    assert target.enjeux.count() == 0


def test_import_missing_file_returns_400():
    user = SuperAdminFactory()
    target = PlanGestionFactory(id_utilisateur_ajout=user)
    client = APIClient()
    client.force_authenticate(user=user)

    resp = client.post(_url(target, "import-arborescence/"), {}, format="multipart")
    assert resp.status_code == 400


# ---------------------------------------------------------------------------
# Import réel + verrou brouillon
# ---------------------------------------------------------------------------


def test_import_creates_arborescence_on_draft():
    user = SuperAdminFactory()
    _base_nomenclatures()
    source = _build_source_plan(user)
    target = PlanGestionFactory(id_utilisateur_ajout=user)

    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(
        _url(target, "import-arborescence/"),
        {"file": _upload_from_plan(source)},
        format="multipart",
    )
    assert resp.status_code == 201, resp.content
    body = resp.json()
    assert body["created"]["enjeux"] == source.enjeux.count()
    assert target.enjeux.count() == source.enjeux.count()


def test_import_refused_on_validated_plan():
    """Verrou #248 : import bloqué (403) hors brouillon, même pour un super admin."""
    user = SuperAdminFactory()
    _base_nomenclatures()
    source = _build_source_plan(user)
    target = PlanGestionValideFactory(id_utilisateur_ajout=user)

    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(
        _url(target, "import-arborescence/"),
        {"file": _upload_from_plan(source)},
        format="multipart",
    )
    assert resp.status_code == 403
    # Aucune écriture sur le plan validé.
    assert target.enjeux.count() == 0


# ---------------------------------------------------------------------------
# Exemples téléchargeables (indépendants d'un plan)
# ---------------------------------------------------------------------------


def test_example_arborescence_endpoint_ok():
    user = SuperAdminFactory()
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.get("/api/plans/plans/example-arborescence-xlsx/")
    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME
    assert resp.content[:2] == b"PK"


def test_example_actions_endpoint_ok():
    user = SuperAdminFactory()
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.get("/api/plans/plans/example-actions-xlsx/")
    assert resp.status_code == 200
    assert resp["Content-Type"] == XLSX_MIME


def test_example_endpoints_require_auth():
    assert APIClient().get(
        "/api/plans/plans/example-arborescence-xlsx/"
    ).status_code in (401, 403)
    assert APIClient().get("/api/plans/plans/example-actions-xlsx/").status_code in (
        401,
        403,
    )


# ---------------------------------------------------------------------------
# Édition assistée (#9) : schéma, validate-data, import-data, read-xlsx (#10)
# ---------------------------------------------------------------------------


def test_import_schema_endpoint():
    client = APIClient()
    client.force_authenticate(user=SuperAdminFactory())
    resp = client.get("/api/plans/plans/import-arborescence-schema/")
    assert resp.status_code == 200
    sheets = resp.json()["sheets"]
    assert any(s["key"] == "enjeux" for s in sheets)
    enjeux = next(s for s in sheets if s["key"] == "enjeux")
    assert any(c["key"] == "code" and c["required"] for c in enjeux["columns"])


def test_validate_data_roundtrip():
    user = SuperAdminFactory()
    _base_nomenclatures()
    target = PlanGestionFactory(id_utilisateur_ajout=user)
    client = APIClient()
    client.force_authenticate(user=user)
    body = {
        "data": {
            "enjeux": [
                {"code": "E1", "categorie": "ENJEU", "libelle": "Qualité des eaux"}
            ]
        }
    }
    resp = client.post(
        f"/api/plans/plans/{target.pk}/import-arborescence/validate-data/",
        body,
        format="json",
    )
    assert resp.status_code == 200
    j = resp.json()
    assert j["can_import"] is True
    assert len(j["data"]["enjeux"]) == 1
    assert target.enjeux.count() == 0  # dry-run


def test_import_data_creates_arborescence():
    user = SuperAdminFactory()
    _base_nomenclatures()
    target = PlanGestionFactory(id_utilisateur_ajout=user)
    client = APIClient()
    client.force_authenticate(user=user)
    body = {
        "data": {
            "enjeux": [
                {"code": "E1", "categorie": "ENJEU", "libelle": "Qualité des eaux"}
            ]
        }
    }
    resp = client.post(
        f"/api/plans/plans/{target.pk}/import-arborescence/import-data/",
        body,
        format="json",
    )
    assert resp.status_code == 201, resp.content
    assert target.enjeux.count() == 1


def test_import_data_refused_on_validated_plan():
    user = SuperAdminFactory()
    _base_nomenclatures()
    target = PlanGestionValideFactory(id_utilisateur_ajout=user)
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(
        f"/api/plans/plans/{target.pk}/import-arborescence/import-data/",
        {"data": {"enjeux": [{"code": "E1", "categorie": "ENJEU", "libelle": "X"}]}},
        format="json",
    )
    assert resp.status_code == 403


def test_read_xlsx_endpoint():
    user = SuperAdminFactory()
    client = APIClient()
    client.force_authenticate(user=user)
    content = build_arborescence_workbook(plan=None)
    upload = SimpleUploadedFile("foreign.xlsx", content, content_type=XLSX_MIME)
    resp = client.post(
        "/api/plans/plans/read-xlsx/", {"file": upload}, format="multipart"
    )
    assert resp.status_code == 200
    sheets = resp.json()["sheets"]
    assert any(s["name"] == "Enjeux" for s in sheets)


# ---------------------------------------------------------------------------
# Modes d'import : add / replace + résumé du contenu existant
# ---------------------------------------------------------------------------


def test_existing_summary_endpoint():
    from tests.factories.enjeux import EnjeuFactory

    user = SuperAdminFactory()
    plan = PlanGestionFactory(id_utilisateur_ajout=user)
    EnjeuFactory(id_pg=plan, id_utilisateur_ajout=user)
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.get(
        f"/api/plans/plans/{plan.pk}/import-arborescence/existing-summary/"
    )
    assert resp.status_code == 200
    assert resp.json()["enjeux"] == 1


def test_import_data_add_mode_appends():
    from tests.factories.enjeux import EnjeuFactory

    user = SuperAdminFactory()
    _base_nomenclatures()
    plan = PlanGestionFactory(id_utilisateur_ajout=user)
    EnjeuFactory(id_pg=plan, id_utilisateur_ajout=user, libelle="Existant")
    client = APIClient()
    client.force_authenticate(user=user)
    body = {
        "mode": "add",
        "data": {
            "enjeux": [{"code": "E1", "categorie": "ENJEU", "libelle": "Nouveau"}]
        },
    }
    resp = client.post(
        f"/api/plans/plans/{plan.pk}/import-arborescence/import-data/",
        body,
        format="json",
    )
    assert resp.status_code == 201, resp.content
    assert plan.enjeux.count() == 2


def test_import_data_create_mode_refuses_non_empty():
    from tests.factories.enjeux import EnjeuFactory

    user = SuperAdminFactory()
    _base_nomenclatures()
    plan = PlanGestionFactory(id_utilisateur_ajout=user)
    EnjeuFactory(id_pg=plan, id_utilisateur_ajout=user)
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(
        f"/api/plans/plans/{plan.pk}/import-arborescence/import-data/",
        {"data": {"enjeux": [{"code": "E1", "categorie": "ENJEU", "libelle": "X"}]}},
        format="json",
    )
    assert resp.status_code == 400  # création seule refuse un plan non vide


# ---------------------------------------------------------------------------
# Import des actions (module 2) : validation, import, verrou brouillon
# ---------------------------------------------------------------------------


def _actions_upload(plan, fill):
    """Classeur d'actions du plan, rempli par ``fill(workbook)``."""
    import io

    from openpyxl import load_workbook

    from apps.plans.services_import_actions import build_actions_workbook

    wb = load_workbook(io.BytesIO(build_actions_workbook(plan)))
    fill(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("actions.xlsx", buf.getvalue(), content_type=XLSX_MIME)


def _write(ws, values):
    """Écrit la ligne 3 (à la place de la ligne exemple) par en-tête."""
    headers = {c.value: c.column for c in ws[2]}
    for col in headers.values():
        ws.cell(row=3, column=col).value = None
    for header, value in values.items():
        ws.cell(row=3, column=headers[header], value=value)


def _plan_for_actions(user):
    from tests.apps.plans.test_import_actions import _plan_with_indicateur

    plan, _ind = _plan_with_indicateur(user)
    return plan


def _fill_ok(wb):
    _write(wb["Actions"], {
        "code": "A1", "indicateur": "I1", "libellé": "Fauche",
        "année début": 2024, "année fin": 2025,
        "mode de ventilation": "Par type de budget",
        "déclinaison par type de coût": "Non",
    })
    _write(wb["Budgets"], {
        "action": "A1", "année": 2024,
        "budget fonctionnement (€)": 800, "budget investissement (€)": 200,
    })
    _write(wb["RH"], {
        "action": "A1", "année": 2024, "jours": 4,
        "catégorie de dépense": "Investissement",
    })


def _fill_wrong_column(wb):
    _fill_ok(wb)
    _write(wb["Budgets"], {"action": "A1", "année": 2024, "budget total (€)": 900})


def test_import_actions_validate_returns_mode_errors():
    """Le rapport situe l'erreur (onglet, ligne, colonne) et nomme le mode."""
    user = SuperAdminFactory()
    plan = _plan_for_actions(user)
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(
        _url(plan, "import-actions/validate/"),
        {"file": _actions_upload(plan, _fill_wrong_column)},
        format="multipart",
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["can_import"] is False
    [issue] = [i for i in body["issues"] if i["level"] == "error"]
    assert (issue["sheet"], issue["row"], issue["column"]) == (
        "Budgets", 3, "budget_total",
    )
    assert "Par type de budget" in issue["message"]


def test_import_actions_creates_with_budget_and_rh():
    from decimal import Decimal

    from apps.plans.models_operations import Operation, OperationAnnee

    user = SuperAdminFactory()
    plan = _plan_for_actions(user)
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(
        _url(plan, "import-actions/"),
        {"file": _actions_upload(plan, _fill_ok)},
        format="multipart",
    )
    assert resp.status_code == 201, resp.json()
    assert resp.json()["created"] == {"actions": 1, "annees": 2, "budgets": 1, "rh": 1}
    op = Operation.objects.get(libelle="Fauche")
    assert (op.ventilation_mode, op.declinaison_par_type_cout) == ("by_type", False)
    oa = OperationAnnee.objects.get(id_operation=op, annee=2024)
    assert oa.budget == Decimal("1000")
    assert oa.rh_lignes.get().categorie_depense == "investissement"


def test_import_actions_invalid_returns_400_with_report():
    from apps.plans.models_operations import Operation

    user = SuperAdminFactory()
    plan = _plan_for_actions(user)
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(
        _url(plan, "import-actions/"),
        {"file": _actions_upload(plan, _fill_wrong_column)},
        format="multipart",
    )
    assert resp.status_code == 400
    assert resp.json()["can_import"] is False
    assert Operation.objects.count() == 0  # transaction : rien d'écrit


def test_import_actions_refused_on_validated_plan():
    from apps.plans.models_operations import Operation

    user = SuperAdminFactory()
    plan = _plan_for_actions(user)
    upload = _actions_upload(plan, _fill_ok)
    plan.statut = "valide"
    plan.save(update_fields=["statut"])
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(_url(plan, "import-actions/"), {"file": upload}, format="multipart")
    assert resp.status_code == 403
    assert Operation.objects.count() == 0


def test_import_actions_missing_file_returns_400():
    user = SuperAdminFactory()
    plan = _plan_for_actions(user)
    client = APIClient()
    client.force_authenticate(user=user)
    resp = client.post(_url(plan, "import-actions/validate/"), {}, format="multipart")
    assert resp.status_code == 400
