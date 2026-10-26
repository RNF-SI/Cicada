"""
#754 — exports budget suivi et RH suivi : feuille « Détail par action ».

Une ligne par poste de dépense (budget) ou par cible de temps (RH) de chaque
action, rattachée à son enjeu, son OLT et son OO, avec le prévu et le réalisé
de chaque année et le total — puis une ligne « Total action ».
"""

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.plans.models_operations import (
    OperationAnnee, OperationAnneeRH, Poste, RealisationOperationAnnee,
    RealisationOperationAnneeRH,
)
from apps.plans.services_export_budget_rh import (
    build_budget_previsionnel_workbook, build_budget_suivi_workbook,
    build_rh_previsionnel_workbook, build_rh_suivi_workbook,
)
from tests.factories.enjeux import (
    EnjeuFactory, IndicateurFactory, MetriqueFactory, NiveauExigenceFactory,
    NomenclatureTypeIndicateurFactory, ObjectifLongTermeFactory,
    ObjectifOperationnelFactory, OperationFactory, ResultatAttenduFactory,
)
from tests.factories.plans import PlanGestionFactory
from tests.factories.users import OrganismeFactory

SHEET = "Détail par action"


@pytest.fixture
def plan(db):
    """Action « type de budget + poste » rattachée à un OLT (indicateur d'état)
    et à un OO (indicateur de réponse). 2024 : 10 j prévus à 300 €/j + 500 € de
    prestataire ; 8 j réalisés + 400 € de prestataire réalisé."""
    p = PlanGestionFactory(annee_debut=2024, annee_fin=2025, statut='valide')
    enjeu = EnjeuFactory(id_pg=p, libelle="Zones humides")
    olt = ObjectifLongTermeFactory(id_enjeu=enjeu, libelle="Maintenir les prairies humides")
    etat = NomenclatureTypeIndicateurFactory(cd_nomenclature='ETAT', mnemonique='ETAT', label='État')
    reponse = NomenclatureTypeIndicateurFactory(
        cd_nomenclature='REPONSE', mnemonique='REPONSE', label='Réponse')
    met_etat = MetriqueFactory(id_indicateur=IndicateurFactory(
        id_ne=NiveauExigenceFactory(id_olt=olt), type_indicateur=etat))
    oo = ObjectifOperationnelFactory(id_enjeu=enjeu, libelle="Restaurer la fauche tardive")
    met_rep = MetriqueFactory(id_indicateur=IndicateurFactory(
        id_ne=None, id_resultat_attendu=ResultatAttenduFactory(id_oo=oo), type_indicateur=reponse))
    op = OperationFactory(
        metriques=[met_etat, met_rep], libelle="Fauche tardive",
        ventilation_mode='by_type_poste', declinaison_par_poste=True,
        declinaison_par_type_cout=True, cout_salarial_auto=True,
        annee_min=2024, annee_max=2024,
    )
    poste = Poste.objects.create(id_pg=p, id_organisme=OrganismeFactory(), nombre=1,
                                 cout_jour=300, nom_local="Garde technicien")
    oa = OperationAnnee.objects.create(id_operation=op, annee=2024, periodicite=True,
                                       cout_prestataire=Decimal('500'))
    OperationAnneeRH.objects.create(id_operation_annee=oa, id_poste=poste, jours=10,
                                    categorie_depense='fonctionnement')
    real = RealisationOperationAnnee.objects.create(
        id_operation_annee=oa, cout_prestataire_realise=Decimal('400'))
    RealisationOperationAnneeRH.objects.create(
        id_realisation_operation_annee=real, id_poste=poste, jours=8,
        categorie_depense='fonctionnement')
    return p


def _rows(data: bytes):
    wb = load_workbook(io.BytesIO(data))
    assert SHEET in wb.sheetnames, wb.sheetnames
    ws = wb[SHEET]
    return [[c.value for c in row] for row in ws.iter_rows()]


def _row(rows, label):
    found = [r for r in rows if label in r]
    assert found, f"ligne « {label} » absente"
    return found[0]


@pytest.mark.django_db
class TestBudgetSuivi:
    def test_feuille_detaillee(self, plan):
        rows = _rows(build_budget_suivi_workbook(plan))
        entete = next(r for r in rows if "Poste de dépense" in r)
        for col in ("Enjeu / FCR", "OLT", "OO", "Code action PG", "Libellé de l'action"):
            assert col in entete
        salarial = _row(rows, "Coût salarial — fonctionnement")
        assert "Maintenir les prairies humides" in salarial
        assert "Restaurer la fauche tardive" in salarial
        assert "Fauche tardive" in salarial
        # 2024 : prévu 10 × 300 = 3 000, réalisé 8 × 300 = 2 400.
        assert "3 000" in salarial and "2 400" in salarial
        presta = _row(rows, "Coût prestataire — fonctionnement")
        assert "500" in presta and "400" in presta
        total = _row(rows, "Total action")
        assert "3 500" in total and "2 800" in total

    def test_pas_dans_le_previsionnel(self, plan):
        wb = load_workbook(io.BytesIO(build_budget_previsionnel_workbook(plan)))
        assert SHEET not in wb.sheetnames


@pytest.mark.django_db
class TestRhSuivi:
    def test_feuille_detaillee(self, plan):
        rows = _rows(build_rh_suivi_workbook(plan))
        garde = _row(rows, "Garde technicien")
        assert "Maintenir les prairies humides" in garde
        # Format des jours commun aux feuilles RH (« 10 » ou « 10.00 »).
        nombres = [float(c) for c in garde if isinstance(c, str) and c.replace('.', '', 1).isdigit()]
        assert 10.0 in nombres and 8.0 in nombres
        assert "Total action" in [c for r in rows for c in r]

    def test_pas_dans_le_previsionnel(self, plan):
        wb = load_workbook(io.BytesIO(build_rh_previsionnel_workbook(plan)))
        assert SHEET not in wb.sheetnames


SYNTHESE = "Synthèse par action"


def _synthese(data: bytes):
    wb = load_workbook(io.BytesIO(data))
    assert SYNTHESE in wb.sheetnames, wb.sheetnames
    return [[c.value for c in row] for row in wb[SYNTHESE].iter_rows()]


@pytest.mark.django_db
class TestSyntheseParAction:
    """#754 (retour 23/10) — une ligne par action, budget total (fonctionnement
    + investissement) ou temps total, prévu / réalisé par année et TOTAL."""

    def test_budget_suivi(self, plan):
        OperationAnnee.objects.filter(id_operation__libelle="Fauche tardive").update(
            autre_cout_invest=Decimal('200'))
        rows = _synthese(build_budget_suivi_workbook(plan))
        entete = next(r for r in rows if "Libellé de l'action" in r)
        for col in ("Enjeu / FCR", "OLT", "OO", "Code simplifié", "Code action PG"):
            assert col in entete
        assert "Poste de dépense" not in entete
        actions = [r for r in rows if "Fauche tardive" in r]
        assert len(actions) == 1   # une seule ligne par action
        ligne = actions[0]
        assert "Maintenir les prairies humides" in ligne
        # 3 000 salarial + 500 prestataire + 200 investissement ; réalisé 2 400 + 400.
        assert "3 700" in ligne and "2 800" in ligne
        assert any(r[0] == "TOTAL" for r in rows)

    def test_rh_suivi(self, plan):
        rows = _synthese(build_rh_suivi_workbook(plan))
        ligne = next(r for r in rows if "Fauche tardive" in r)
        nombres = [float(c) for c in ligne if isinstance(c, str) and c.replace('.', '', 1).isdigit()]
        assert 10.0 in nombres and 8.0 in nombres

    def test_pas_dans_les_previsionnels(self, plan):
        for build in (build_budget_previsionnel_workbook, build_rh_previsionnel_workbook):
            assert SYNTHESE not in load_workbook(io.BytesIO(build(plan))).sheetnames
