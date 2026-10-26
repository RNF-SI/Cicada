"""
#735 (retour 21/10) — les exports budget / RH du plan (prévisionnel et suivi)
suivent la ventilation la plus détaillée de ses actions, et se simplifient
sinon : pas d'onglet par organisme si aucune action n'est ventilée par
organisme, pas de fonctionnement / investissement si aucune ne ventile par type
de budget, pas de synthèse par poste de dépense sans déclinaison par type de
coût, pas de synthèse par type de poste sans temps décliné par poste.

Les totaux ne changent jamais : seule la présentation s'adapte.
"""

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.plans.models_operations import (
    Fonction, OperationAnnee, OperationAnneeOrganisme, OperationAnneeRH, Poste,
    PosteFonction, RealisationOperationAnnee,
)
from apps.plans.services_export_budget_rh import (
    build_budget_previsionnel_workbook, build_budget_suivi_workbook,
    build_rh_previsionnel_workbook, build_rh_suivi_workbook,
)
from tests.factories.enjeux import (
    EnjeuFactory, IndicateurFactory, MetriqueFactory, NiveauExigenceFactory,
    ObjectifLongTermeFactory, OperationFactory,
)
from tests.factories.plans import PlanGestionFactory
from tests.factories.users import OrganismeFactory

pytestmark = pytest.mark.django_db

BUILDERS = [
    build_budget_previsionnel_workbook, build_budget_suivi_workbook,
    build_rh_previsionnel_workbook, build_rh_suivi_workbook,
]


def _wb(builder, plan):
    return load_workbook(io.BytesIO(builder(plan)))


def _texts(wb):
    """Tout le texte des classeurs, onglets compris (minuscules)."""
    out = [s.lower() for s in wb.sheetnames]
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            out += [str(v).lower() for v in row if isinstance(v, str)]
    return " | ".join(out)


def _total_cell(ws, label, real=False):
    """Valeur de la colonne TOTAL (prévu, ou réalisé) de la ligne `label`."""
    hdr = next(r for r in range(1, ws.max_row + 1)
               if any(ws.cell(r, c).value == "TOTAL" for c in range(1, ws.max_column + 1)))
    col = next(c for c in range(1, ws.max_column + 1) if ws.cell(hdr, c).value == "TOTAL")
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == label)
    return ws.cell(row, col + (1 if real else 0)).value


@pytest.fixture
def plan():
    p = PlanGestionFactory(annee_debut=2024, annee_fin=2025)
    ne = NiveauExigenceFactory(id_olt=ObjectifLongTermeFactory(id_enjeu=EnjeuFactory(id_pg=p)))
    p._met = MetriqueFactory(id_indicateur=IndicateurFactory(id_ne=ne))
    return p


def _action(plan, code, **kw):
    return OperationFactory(metriques=[plan._met], code_operation=code,
                            annee_min=2024, annee_max=2025, **kw)


def _none_sans_detail(plan, code="A1"):
    """Action « aucune ventilation », total direct : 1 000 € et 5 j en 2024."""
    op = _action(plan, code, ventilation_mode="none", declinaison_par_type_cout=False)
    oa = OperationAnnee.objects.create(id_operation=op, annee=2024, periodicite=True,
                                       budget=Decimal("1000"))
    OperationAnneeRH.objects.create(id_operation_annee=oa, jours=5,
                                    categorie_depense="fonctionnement")
    RealisationOperationAnnee.objects.create(id_operation_annee=oa, budget_realise=Decimal("800"))
    return op


def _max(plan, code="B1"):
    """Action ventilation maximale : poste de l'organisme, invest + fonct."""
    org = OrganismeFactory(nom_organisme="Org Beta")
    fonction, _ = Fonction.objects.get_or_create(
        libelle="Technicien (export ventilation)", defaults={"type_poste": "salarie"})
    poste = Poste.objects.create(id_pg=plan, id_organisme=org, nombre=1, cout_jour=200)
    PosteFonction.objects.create(id_poste=poste, id_fonction=fonction, pourcentage=100)
    op = _action(plan, code, ventilation_mode="by_org_type_poste",
                 declinaison_par_poste=True, declinaison_par_type_cout=True)
    oa = OperationAnnee.objects.create(id_operation=op, annee=2024, periodicite=True)
    OperationAnneeRH.objects.create(id_operation_annee=oa, id_poste=poste, jours=3,
                                    categorie_depense="fonctionnement")
    OperationAnneeOrganisme.objects.create(id_operation_annee=oa, id_organisme=org,
                                           cout_stage=40, cout_prestataire_invest=60)
    return op


def _by_org(plan, code="C1"):
    """Action « par organisme » sans type de budget : 300 € pour l'organisme."""
    org = OrganismeFactory(nom_organisme="Org Gamma")
    op = _action(plan, code, ventilation_mode="by_org", declinaison_par_type_cout=False)
    oa = OperationAnnee.objects.create(id_operation=op, annee=2024, periodicite=True)
    OperationAnneeOrganisme.objects.create(id_operation_annee=oa, id_organisme=org,
                                           budget_fonctionnement=300)
    OperationAnneeRH.objects.create(id_operation_annee=oa, id_organisme=org, jours=2,
                                    categorie_depense="fonctionnement")
    return op


class TestPlanSansVentilation:
    """Toutes les actions en « aucune », sans déclinaison : export minimal."""

    @pytest.mark.parametrize("builder", BUILDERS)
    def test_ni_fonctionnement_ni_investissement_ni_organisme(self, plan, builder):
        _none_sans_detail(plan)
        wb = _wb(builder, plan)
        texte = _texts(wb)
        assert "fonctionnement" not in texte
        assert "investissement" not in texte
        assert "non ventilé" not in texte
        # #754 — les exports de SUIVI ajoutent la synthèse puis le détail par
        # action (retour 23/10 : une ligne par action).
        suivi = builder in (build_budget_suivi_workbook, build_rh_suivi_workbook)
        assert wb.sheetnames == (["TOTAL", "Synthèse par action", "Détail par action"]
                                 if suivi else ["TOTAL"])

    def test_totaux_inchanges(self, plan):
        _none_sans_detail(plan)
        ws = _wb(build_budget_suivi_workbook, plan)["TOTAL"]
        assert _total_cell(ws, "TOTAL — Budget") == "1 000"
        assert _total_cell(ws, "TOTAL — Budget", real=True) == "800"
        ws = _wb(build_rh_previsionnel_workbook, plan)["TOTAL"]
        assert _total_cell(ws, "TOTAL jours prévus") in ("5", "5.00")


class TestPlanMixte:
    """Une action « aucune » + une action en ventilation maximale : la plus
    détaillée l'emporte pour tout le plan."""

    def test_budget_garde_organismes_types_et_detail(self, plan):
        _none_sans_detail(plan)
        _max(plan)
        wb = _wb(build_budget_previsionnel_workbook, plan)
        assert any("Org Beta" in s for s in wb.sheetnames)
        assert "TOTAL par poste de dépense" in wb.sheetnames
        texte = _texts(wb)
        assert "fonctionnement" in texte and "investissement" in texte
        # Stage détaillé dans la synthèse ; totaux inchangés : 1000 + 3×200 + 40 + 60.
        ws = wb["TOTAL par poste de dépense"]
        assert any(ws.cell(r, 1).value == "Coût stage" for r in range(1, ws.max_row + 1))
        assert _total_cell(ws, "TOTAL") == "1 700"

    def test_rh_garde_la_synthese_par_type_de_poste(self, plan):
        _none_sans_detail(plan)
        _max(plan)
        assert "RH par type de poste" in _wb(build_rh_previsionnel_workbook, plan).sheetnames
        assert "Total par poste" in _wb(build_rh_suivi_workbook, plan).sheetnames


class TestPlanParOrganismeSansType:
    """Ventilation par organisme seule : onglets par organisme, sans
    fonctionnement / investissement ni synthèse de dépense."""

    @pytest.mark.parametrize("builder", BUILDERS)
    def test_organismes_sans_type_de_budget(self, plan, builder):
        _by_org(plan)
        wb = _wb(builder, plan)
        assert any("Org Gamma" in s for s in wb.sheetnames)
        texte = _texts(wb)
        assert "fonctionnement" not in texte
        assert "investissement" not in texte
        assert "par poste" not in texte

    def test_totaux_inchanges(self, plan):
        _by_org(plan)
        wb = _wb(build_budget_previsionnel_workbook, plan)
        ws = next(wb[s] for s in wb.sheetnames if "Org Gamma" in s)
        assert _total_cell(ws, "TOTAL — Budget Org Gamma") == "300"
        assert _total_cell(wb["TOTAL OG"], "TOTAL — Budget") == "300"


class TestDetailSansTypeDeBudget:
    """Déclinaison par type de coût sans type de budget : synthèse par type de
    coût (salarial, stage, prestataire, autres), sans fonct / invest."""

    def test_synthese_par_type_de_cout(self, plan):
        op = _action(plan, "D1", ventilation_mode="none", declinaison_par_type_cout=True,
                     cout_salarial_auto=False)
        oa = OperationAnnee.objects.create(id_operation=op, annee=2024, periodicite=True,
                                           cout_salarial=100, cout_stage=20,
                                           cout_prestataire=30, autre_cout=10)
        wb = _wb(build_budget_previsionnel_workbook, plan)
        ws = wb["TOTAL par poste de dépense"]
        labels = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
        for lab in ("► Coût salarial", "Coût stage", "Prestataire", "Autres coûts"):
            assert lab in labels, lab
        assert _total_cell(ws, "Coût stage") == "20"
        assert _total_cell(ws, "Autres coûts") == "10"
        assert _total_cell(ws, "TOTAL") == "160"
        assert "fonctionnement" not in _texts(wb)
        assert oa.pk
