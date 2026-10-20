"""
#735 — programmation GLOBALE d'une action et déclinaison par type de coût en
modes « aucune ventilation » / « par organisme ».

Une action programmée globalement porte son budget et son temps de travail sur
une seule ligne ``OperationAnnee`` sans année ; ses années ne gardent que leur
programmation. Cette ligne ne doit apparaître dans aucune lecture « par année »,
mais compter dans tous les totaux.
"""

import io
from collections import defaultdict
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.plans.models_operations import (
    OperationAnnee, OperationAnneeOrganisme, OperationAnneeRH,
)
from apps.plans.services_export_finance import (
    GLOBAL_YEAR, GLOBAL_YEAR_LABEL, build_action_finance, poste_entry_factory,
)
from tests.factories.enjeux import (
    EnjeuFactory, IndicateurFactory, MetriqueFactory, NiveauExigenceFactory,
    ObjectifLongTermeFactory, OperationFactory,
)
from tests.factories.plans import PlanGestionFactory
from tests.factories.users import OrganismeFactory, SuperAdminFactory


@pytest.fixture
def action_globale(db):
    """Action 2028-2029 « par type de budget », 20 000 € et 150 j sur la période."""
    plan = PlanGestionFactory(annee_debut=2027, annee_fin=2030, statut='draft')
    enjeu = EnjeuFactory(id_pg=plan)
    ne = NiveauExigenceFactory(id_olt=ObjectifLongTermeFactory(id_enjeu=enjeu))
    met = MetriqueFactory(id_indicateur=IndicateurFactory(id_ne=ne))
    op = OperationFactory(
        metriques=[met], ventilation_mode='by_type', declinaison_par_type_cout=False,
        programmation_globale=True, annee_min=2028, annee_max=2029,
    )
    for annee in (2028, 2029):
        OperationAnnee.objects.create(id_operation=op, annee=annee, periodicite=True)
    globale = OperationAnnee.objects.create(
        id_operation=op, annee=None,
        budget_fonctionnement=Decimal('15000'), budget_investissement=Decimal('5000'),
    )
    OperationAnneeRH.objects.create(
        id_operation_annee=globale, jours=150, categorie_depense='fonctionnement')
    return {'plan': plan, 'op': op, 'globale': globale}


@pytest.mark.django_db
class TestLigneGlobale:
    def test_absente_des_lectures_annuelles(self, action_globale):
        op = action_globale['op']
        assert [oa.annee for oa in op.operation_annees.all()] == [2028, 2029]
        assert OperationAnnee.objects.filter(id_operation=op).count() == 2
        assert op.annee_globale == action_globale['globale']

    def test_supprimee_avec_l_action(self, action_globale):
        op = action_globale['op']
        op.delete()
        assert not OperationAnnee.tous.filter(annee__isnull=True).exists()


@pytest.mark.django_db
class TestApiProgrammationGlobale:
    def _client(self):
        client = APIClient()
        client.force_authenticate(SuperAdminFactory())
        return client

    def test_lecture(self, action_globale):
        op = action_globale['op']
        resp = self._client().get(f'/api/plans/operations/{op.pk}/')
        assert resp.status_code == 200
        data = resp.json()
        assert data['programmation_globale'] is True
        assert [a['annee'] for a in data['operation_annees']] == [2028, 2029]
        assert data['operation_annee_globale']['annee'] is None
        assert Decimal(data['operation_annee_globale']['budget_fonctionnement']) == 15000
        assert data['operation_annee_globale']['rh_lignes'][0]['jours'] == '150.00'

    def test_creation_sans_case_garde_le_total(self, action_globale):
        """#735 — sans « déclinaison » explicite, le mode « aucune » reste en total."""
        met = action_globale['op'].metriques.first()
        resp = self._client().post('/api/plans/operations/', {
            'libelle': 'API sans case', 'ventilation_mode': 'none',
            'metrique_ids': [met.id_metrique],
        }, format='json')
        assert resp.status_code == 201, resp.content
        assert resp.json()['declinaison_par_type_cout'] is False

    def test_ecriture_puis_retour_a_l_annuel(self, action_globale):
        op = action_globale['op']
        client = self._client()
        annees = [{'annee': a, 'periodicite': True, 'organismes': [], 'rh_lignes': []}
                  for a in (2028, 2029)]
        resp = client.patch(f'/api/plans/operations/{op.pk}/', {
            'programmation_globale': True,
            'operation_annees': annees,
            'operation_annee_globale': {
                'budget_fonctionnement': '12000', 'budget_investissement': '8000',
                'organismes': [], 'rh_lignes': [{'jours': '90', 'categorie_depense': 'fonctionnement'}],
            },
        }, format='json')
        assert resp.status_code == 200, resp.content
        globale = OperationAnnee.tous.get(id_operation=op, annee__isnull=True)
        assert globale.budget_fonctionnement == Decimal('12000')
        assert list(globale.rh_lignes.values_list('jours', flat=True)) == [Decimal('90')]

        resp = client.patch(f'/api/plans/operations/{op.pk}/', {
            'programmation_globale': False, 'operation_annees': annees,
            'operation_annee_globale': None,
        }, format='json')
        assert resp.status_code == 200, resp.content
        assert not OperationAnnee.tous.filter(id_operation=op, annee__isnull=True).exists()
        assert resp.json()['operation_annee_globale'] is None


@pytest.mark.django_db
class TestExportsProgrammationGlobale:
    def test_prevu_dans_la_colonne_periode(self, action_globale):
        af = build_action_finance(action_globale['op'], {}, defaultdict(poste_entry_factory))
        assert af.year_total(2028).tot == 0
        periode = af.year_total(GLOBAL_YEAR)
        assert periode.tot_fonct == Decimal('15000')
        assert periode.tot_invest == Decimal('5000')
        assert periode.jours == Decimal('150')

    def test_budget_previsionnel_colonne_periode_et_total(self, action_globale):
        from apps.plans.services_export_budget_rh import build_budget_previsionnel_workbook
        wb = load_workbook(io.BytesIO(build_budget_previsionnel_workbook(action_globale['plan'])))
        ws = wb['TOTAL'] if 'TOTAL' in wb.sheetnames else wb[wb.sheetnames[0]]
        values = [[ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
                  for r in range(1, ws.max_row + 1)]
        header = next(row for row in values if 'TOTAL' in row)
        assert GLOBAL_YEAR_LABEL in header
        col = header.index(GLOBAL_YEAR_LABEL)
        periode = [row[col] for row in values if row[col] not in (None, GLOBAL_YEAR_LABEL)]
        # Prévu de la période (fonctionnement 15 000 €, investissement 5 000 €).
        assert '15 000' in periode and '5 000' in periode

    def test_fiche_action_colonne_periode(self, action_globale):
        from apps.plans.services_export_fiche_action import build_fiche_action_workbook
        wb = load_workbook(io.BytesIO(build_fiche_action_workbook(action_globale['plan'])))
        ws = wb[wb.sheetnames[0]]
        flat = [ws.cell(r, c).value for r in range(1, ws.max_row + 1)
                for c in range(1, ws.max_column + 1)]
        assert GLOBAL_YEAR_LABEL in flat

    def test_fiche_action_word(self, action_globale):
        from docx import Document
        from apps.plans.services_export_fiche_action_word import build_fiche_action_docx
        doc = Document(io.BytesIO(build_fiche_action_docx(action_globale['plan'])))
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        assert GLOBAL_YEAR_LABEL in cells

    def test_copie_d_action_copie_la_ligne_globale(self, action_globale):
        from apps.plans.element_copy import ElementCopyService
        new_op = ElementCopyService._copy_operation(action_globale['op'], SuperAdminFactory())
        assert new_op.programmation_globale is True
        assert new_op.annee_globale.budget_fonctionnement == Decimal('15000')
        assert new_op.annee_globale.rh_lignes.count() == 1


@pytest.mark.django_db
class TestDetailSansTypeDeBudget:
    """#735 d — « aucune ventilation » / « par organisme » avec détail des coûts."""

    def test_mode_aucune_ne_compte_pas_le_total_en_plus(self, db):
        op = OperationFactory(
            ventilation_mode='none', declinaison_par_type_cout=True, cout_salarial_auto=False)
        OperationAnnee.objects.create(
            id_operation=op, annee=2027, periodicite=True,
            budget=Decimal('1700'), cout_salarial=Decimal('1000'), cout_stage=Decimal('200'),
            cout_prestataire=Decimal('400'), autre_cout=Decimal('100'))
        c = build_action_finance(op, {}, defaultdict(poste_entry_factory)).year_total(2027)
        assert c.sal_fonct == Decimal('1000')
        assert c.tot == Decimal('1700')

    def test_mode_aucune_sans_detail_garde_le_total(self, db):
        op = OperationFactory(ventilation_mode='none', declinaison_par_type_cout=False)
        OperationAnnee.objects.create(
            id_operation=op, annee=2027, periodicite=True, budget=Decimal('900'))
        c = build_action_finance(op, {}, defaultdict(poste_entry_factory)).year_total(2027)
        assert c.tot == Decimal('900')

    def test_mode_par_organisme_detail(self, db):
        op = OperationFactory(
            ventilation_mode='by_org', declinaison_par_type_cout=True, cout_salarial_auto=False)
        oa = OperationAnnee.objects.create(id_operation=op, annee=2027, periodicite=True)
        org = OrganismeFactory()
        OperationAnneeOrganisme.objects.create(
            id_operation_annee=oa, id_organisme=org,
            cout_salarial=Decimal('500'), cout_prestataire=Decimal('300'))
        af = build_action_finance(op, {}, defaultdict(poste_entry_factory))
        assert af.cell(org.id_organisme, 2027).tot == Decimal('800')
