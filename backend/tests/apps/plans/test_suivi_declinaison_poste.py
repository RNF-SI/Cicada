"""
#760 — case « Déclinaison par type de poste » du SUIVI (temps de travail réalisé).

- Décochée : temps réalisé saisi au global (ou par organisme en ventilation par
  organisme), sans choix de poste.
- Cochée : saisie par poste ; chaque ligne de temps doit alors viser un poste.
- Par défaut : cochée si la programmation est déclinée par poste — et alors
  non décochable (on ne suit pas moins finement qu'on n'a prévu) —, décochée
  sinon, « aucune ventilation » comprise.
"""

from decimal import Decimal

import pytest
from rest_framework.test import APIClient

from apps.plans.models_operations import (
    OperationAnnee, OperationRealisationGlobale, Poste, RealisationOperationAnnee,
)
from tests.factories.enjeux import (
    EnjeuFactory, IndicateurFactory, MetriqueFactory, NiveauExigenceFactory,
    ObjectifLongTermeFactory, OperationFactory,
)
from tests.factories.plans import PlanGestionFactory
from tests.factories.users import OrganismeFactory, SuperAdminFactory


@pytest.fixture
def plan(db):
    p = PlanGestionFactory(annee_debut=2026, annee_fin=2028, statut='valide')
    ne = NiveauExigenceFactory(id_olt=ObjectifLongTermeFactory(id_enjeu=EnjeuFactory(id_pg=p)))
    p._met = MetriqueFactory(id_indicateur=IndicateurFactory(id_ne=ne))
    p._poste = Poste.objects.create(id_pg=p, id_organisme=OrganismeFactory(), nombre=1, cout_jour=300)
    return p


def _action(plan, **kw):
    op = OperationFactory(metriques=[plan._met], **kw)
    oa = OperationAnnee.objects.create(id_operation=op, annee=2027, periodicite=True)
    return op, oa


def _sans_poste(plan):
    return _action(plan, ventilation_mode='none', declinaison_par_type_cout=False)


def _par_poste(plan):
    return _action(plan, ventilation_mode='by_type_poste', declinaison_par_poste=True,
                   declinaison_par_type_cout=True)


def _client():
    client = APIClient()
    client.force_authenticate(SuperAdminFactory())
    return client


def _reglages(op):
    return f'/api/plans/realisations/suivi-reglages/{op.pk}/'


@pytest.mark.django_db
class TestValeurParDefaut:
    def test_decochee_sans_declinaison_a_la_programmation(self, plan):
        op, _ = _sans_poste(plan)
        assert op.get_declinaison_par_poste_realise() is False
        data = _client().get(_reglages(op)).json()
        assert data['declinaison_par_poste_realise'] is False
        assert data['declinaison_poste_verrouillee'] is False
        assert _client().get(f'/api/plans/operations/{op.pk}/').json()[
            'declinaison_par_poste_realise'] is False

    def test_cochee_et_verrouillee_si_la_programmation_est_par_poste(self, plan):
        op, _ = _par_poste(plan)
        data = _client().get(_reglages(op)).json()
        assert data['declinaison_par_poste_realise'] is True
        assert data['declinaison_poste_verrouillee'] is True


@pytest.mark.django_db
class TestReglage:
    def test_cocher_au_suivi_sans_declinaison_a_la_programmation(self, plan):
        op, _ = _sans_poste(plan)
        r = _client().post(_reglages(op), {'declinaison_par_poste_realise': True}, format='json')
        assert r.status_code == 200, r.content
        assert r.json()['declinaison_par_poste_realise'] is True
        op.refresh_from_db()
        assert op.get_declinaison_par_poste_realise() is True

    def test_non_decochable_si_la_programmation_est_par_poste(self, plan):
        op, _ = _par_poste(plan)
        r = _client().post(_reglages(op), {'declinaison_par_poste_realise': False}, format='json')
        assert r.status_code == 400
        assert 'declinaison_par_poste_realise' in r.json()

    def test_retirer_la_surcharge_de_niveau_garde_le_reglage(self, plan):
        op, _ = _sans_poste(plan)
        OperationRealisationGlobale.objects.create(id_operation=op, declinaison_par_poste_realise=True)
        r = _client().delete(f'/api/plans/realisations/global-realisation/{op.pk}/')
        assert r.status_code == 200, r.content
        assert OperationRealisationGlobale.objects.get(id_operation=op).declinaison_par_poste_realise is True


@pytest.mark.django_db
class TestSaisieParPoste:
    def _upsert(self, oa, rh):
        return _client().post('/api/plans/realisations/upsert/', {
            'id_operation_annee': oa.id_operation_annee, 'rh_lignes': rh,
        }, format='json')

    def test_ligne_sans_poste_refusee(self, plan):
        op, oa = _sans_poste(plan)
        OperationRealisationGlobale.objects.create(id_operation=op, declinaison_par_poste_realise=True)
        r = self._upsert(oa, [{'jours': '5', 'finance': True}])
        assert r.status_code == 400, r.content
        assert 'rh_lignes' in r.json()
        # Refus sans effet de bord : aucune réalisation créée.
        assert not RealisationOperationAnnee.objects.filter(id_operation_annee=oa).exists()

    def test_ligne_par_poste_acceptee(self, plan):
        op, oa = _sans_poste(plan)
        OperationRealisationGlobale.objects.create(id_operation=op, declinaison_par_poste_realise=True)
        r = self._upsert(oa, [{'id_poste': plan._poste.id_poste, 'jours': '5', 'finance': True}])
        assert r.status_code == 200, r.content
        ligne = RealisationOperationAnnee.objects.get(id_operation_annee=oa).rh_lignes.get()
        assert ligne.id_poste_id == plan._poste.id_poste
        assert ligne.jours == Decimal('5')

    def test_annee_sans_temps_acceptee(self, plan):
        """Une année non réalisée n'a pas de temps à attribuer."""
        op, oa = _par_poste(plan)
        assert self._upsert(oa, []).status_code == 200

    def test_decochee_temps_global_accepte(self, plan):
        _op, oa = _sans_poste(plan)
        assert self._upsert(oa, [{'jours': '5', 'finance': True}]).status_code == 200


@pytest.mark.django_db
class TestExportsSuivi:
    """Le réalisé saisi par poste au seul suivi apparaît dans les exports de suivi."""

    def test_rh_suivi_total_par_poste(self, plan):
        import io
        from openpyxl import load_workbook
        from apps.plans.models_operations import RealisationOperationAnneeRH
        from apps.plans.services_export_budget_rh import (
            build_rh_previsionnel_workbook, build_rh_suivi_workbook,
        )

        op, oa = _sans_poste(plan)
        OperationRealisationGlobale.objects.create(id_operation=op, declinaison_par_poste_realise=True)
        real = RealisationOperationAnnee.objects.create(id_operation_annee=oa)
        RealisationOperationAnneeRH.objects.create(
            id_realisation_operation_annee=real, id_poste=plan._poste, jours=7)

        suivi = load_workbook(io.BytesIO(build_rh_suivi_workbook(plan)))
        assert 'Total par poste' in suivi.sheetnames
        # Le prévisionnel, lui, n'est pas décliné par poste.
        prev = load_workbook(io.BytesIO(build_rh_previsionnel_workbook(plan)))
        assert 'RH par type de poste' not in prev.sheetnames


# -----------------------------------------------------------------------------
# #773 — la fiche action a sa propre case « Déclinaison par type de poste » :
# détailler les coûts ne décline plus le temps par poste.
# -----------------------------------------------------------------------------

class TestDeclinaisonPosteProgrammation:
    def test_couts_detailles_sans_case_poste(self, plan):
        op, _ = _action(plan, ventilation_mode='none', declinaison_par_type_cout=True,
                        declinaison_par_poste=False)
        assert op.temps_par_poste() is False
        # Pas de coût jour : le coût salarial réalisé se saisit.
        assert op.get_cout_salarial_auto_realise() is False

    def test_case_poste_hors_mode_poste(self, plan):
        op, _ = _action(plan, ventilation_mode='by_org', declinaison_par_type_cout=False,
                        declinaison_par_poste=True)
        assert op.temps_par_poste() is True
        assert op.get_declinaison_par_poste_realise() is True

    def test_mode_type_poste_force(self, plan):
        op, _ = _action(plan, ventilation_mode='by_type_poste', declinaison_par_poste=False)
        assert op.temps_par_poste() is True

    def test_ventilation_defaults_reprend_la_case(self, plan):
        op, _ = _action(plan, ventilation_mode='none', declinaison_par_type_cout=True,
                        declinaison_par_poste=True)
        r = _client().get(f'/api/plans/operations/ventilation-defaults/{plan.pk}/')
        assert r.status_code == 200
        assert r.data['source_operation_id'] == op.pk
        assert r.data['declinaison_par_poste'] is True
