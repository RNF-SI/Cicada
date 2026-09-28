"""
#671 — Les actions rattachées par la branche « facteur → pression → OO → RA »
doivent figurer dans les exports, au même titre que celles de la branche
« enjeu → OLT → NE ».

Les exports ne suivaient que la seconde : sur le plan « Lacs et zones humides »
du jeu d'essai, 8 actions sur 25 manquaient à l'export des fiches, et les
exports budget / RH sous-estimaient le plan.
"""

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.plans.access import plan_operation_ids
from apps.plans.services_export_fiche_action import build_fiche_action_workbook
from apps.plans.services_export_finance import build_plan_finance
from apps.plans.views_operations import RealisationOperationAnneeViewSet
from tests.factories.enjeux import (
    EnjeuFactory, FacteurInfluenceFactory, IndicateurFactory,
    IndicateurPressionFactory, MetriqueFactory, NiveauExigenceFactory,
    ObjectifLongTermeFactory, ObjectifOperationnelFactory, OperationAnneeFactory,
    OperationFactory, PressionFactory, RealisationOperationAnneeFactory,
    ResultatAttenduFactory,
)
from tests.factories.plans import PlanGestionFactory
from tests.factories.users import SuperAdminFactory


@pytest.mark.django_db
class TestExportsBranchePressions:

    @pytest.fixture
    def plan(self):
        plan = PlanGestionFactory(annee_debut=2024, annee_fin=2028)

        # Branche état : enjeu → OLT → NE → indicateur → métrique → action
        enjeu_etat = EnjeuFactory(id_pg=plan, libelle='Enjeu tourbière')
        ne = NiveauExigenceFactory(id_olt=ObjectifLongTermeFactory(id_enjeu=enjeu_etat))
        op_etat = OperationFactory(code_operation='ETAT1', ventilation_mode='none')
        op_etat.metriques.add(MetriqueFactory(id_indicateur=IndicateurFactory(id_ne=ne)))

        # Branche pression : enjeu → facteur → pression → OO → RA → indicateur
        enjeu_pression = EnjeuFactory(id_pg=plan, libelle='Enjeu qualité de l\'eau')
        pression = PressionFactory(
            id_facteur_influence=FacteurInfluenceFactory(id_enjeu=enjeu_pression))
        oo = ObjectifOperationnelFactory(pressions=[pression])
        ra = ResultatAttenduFactory(id_oo=oo)
        op_pression = OperationFactory(code_operation='PRES1', ventilation_mode='none')
        op_pression.metriques.add(MetriqueFactory(
            id_indicateur=IndicateurPressionFactory(id_resultat_attendu=ra)))
        oa = OperationAnneeFactory(id_operation=op_pression, annee=2025,
                                   budget=Decimal('1200.00'))
        RealisationOperationAnneeFactory(id_operation_annee=oa)

        # Action rattachée directement à l'indicateur (#367), branche pression
        op_direct = OperationFactory(
            code_operation='PRES2',
            id_indicateur=IndicateurPressionFactory(id_resultat_attendu=ra))

        # Bruit : une action d'un autre plan ne doit pas remonter.
        OperationFactory(code_operation='AUTRE').metriques.add(MetriqueFactory())

        return {'plan': plan, 'etat': op_etat, 'pression': op_pression,
                'direct': op_direct, 'enjeu_pression': enjeu_pression}

    def test_toutes_les_branches_sont_retenues(self, plan):
        assert plan_operation_ids(plan['plan']) == {
            plan['etat'].pk, plan['pression'].pk, plan['direct'].pk}

    def test_export_fiches_contient_les_actions_des_pressions(self, plan):
        wb = load_workbook(io.BytesIO(build_fiche_action_workbook(plan['plan'])))
        assert sorted(wb.sheetnames) == ['ETAT1', 'PRES1', 'PRES2']

    def test_cadre_remonte_l_enjeu_par_les_pressions(self, plan):
        wb = load_workbook(io.BytesIO(build_fiche_action_workbook(plan['plan'])))
        ws = wb['PRES1']
        enjeu = next(ws.cell(r, 4).value for r in range(1, ws.max_row + 1)
                     if ws.cell(r, 1).value == 'Enjeu')
        assert enjeu == "Enjeu qualité de l'eau"

    def test_export_d_une_seule_fiche_de_la_branche_pression(self, plan):
        wb = load_workbook(io.BytesIO(build_fiche_action_workbook(
            plan['plan'], operation_ids=[plan['pression'].pk])))
        assert wb.sheetnames == ['PRES1']

    def test_budget_compte_les_actions_des_pressions(self, plan):
        pf = build_plan_finance(plan['plan'])
        actions = {af.libelle: af for af in pf.actions}
        af = actions[plan['pression'].libelle]

        assert len(pf.actions) == 3
        assert af.enjeu_label == "Enjeu qualité de l'eau"
        assert af.year_total(2025).tot == Decimal('1200.00')

    def test_realisations_du_bilan_incluent_les_pressions(self, plan):
        realisations = RealisationOperationAnneeViewSet()._plan_realisations(plan['plan'])
        assert [r.id_operation_annee.id_operation_id for r in realisations] == [
            plan['pression'].pk]

    # --- Par les endpoints, comme l'interface (#671) ---

    @staticmethod
    def _get(url):
        client = APIClient()
        client.force_authenticate(user=SuperAdminFactory())
        resp = client.get(url)
        assert resp.status_code == 200, url
        content = b''.join(resp.streaming_content) if resp.streaming else resp.content
        return load_workbook(io.BytesIO(content))

    def test_endpoint_fiches_actions_du_plan(self, plan):
        wb = self._get(f"/api/plans/plans/{plan['plan'].id_pg}/export-fiches-actions-xlsx/")
        assert sorted(wb.sheetnames) == ['ETAT1', 'PRES1', 'PRES2']

    def test_endpoint_fiche_seule_d_une_action_des_pressions(self, plan):
        """Avant #671 : classeur vide « Ce plan ne contient pas encore d'action »."""
        wb = self._get(f"/api/plans/operations/{plan['pression'].pk}/export-fiche-xlsx/")
        assert wb.sheetnames == ['PRES1']

    def test_endpoint_budget_previsionnel_liste_l_action_des_pressions(self, plan):
        wb = self._get(f"/api/plans/plans/{plan['plan'].id_pg}/export-budget-previsionnel-xlsx/")
        textes = {c.value for ws in wb for row in ws.iter_rows() for c in row
                  if isinstance(c.value, str)}
        assert plan['pression'].libelle in textes
