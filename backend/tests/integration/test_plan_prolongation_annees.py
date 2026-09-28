"""
#672 — Les années ajoutées par la prolongation d'un plan (#250) doivent être
programmables, exportées et suivies.

La prolongation pose `annees_extension` (1 ou 2) sans modifier `annee_fin` :
les exports et le bilan, qui construisaient les années de `annee_debut` à
`annee_fin`, ignoraient donc les années ajoutées.
"""

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.plans.services_export_fiche_action import build_fiche_action_workbook
from apps.plans.services_export_finance import build_plan_finance
from tests.factories.enjeux import (
    EnjeuFactory, IndicateurFactory, MetriqueFactory, NiveauExigenceFactory,
    ObjectifLongTermeFactory, OperationAnneeFactory, OperationFactory,
)
from tests.factories.plans import PlanGestionFactory
from tests.factories.users import SuperAdminFactory


@pytest.mark.django_db
class TestAnneesDuPlan:

    def test_annees_plan_inclut_la_prolongation(self):
        plan = PlanGestionFactory(annee_debut=2021, annee_fin=2025, annees_extension=2)
        assert plan.annee_fin_effective == 2027
        assert plan.annees_plan() == [2021, 2022, 2023, 2024, 2025, 2026, 2027]

    def test_plan_non_prolonge_inchange(self):
        plan = PlanGestionFactory(annee_debut=2021, annee_fin=2025, annees_extension=0)
        assert plan.annee_fin_effective == 2025
        assert plan.annees_plan() == [2021, 2022, 2023, 2024, 2025]

    def test_periode_incomplete(self):
        plan = PlanGestionFactory.build(annee_debut=2021, annee_fin=None)
        assert plan.annee_fin_effective is None
        assert plan.annees_plan() == []


@pytest.mark.django_db
class TestPlanProlongeExportsEtBilan:
    """Une action programmée sur une année de prolongation (2027)."""

    @pytest.fixture
    def plan(self):
        plan = PlanGestionFactory(annee_debut=2021, annee_fin=2025, annees_extension=2)
        ne = NiveauExigenceFactory(
            id_olt=ObjectifLongTermeFactory(id_enjeu=EnjeuFactory(id_pg=plan)))
        op = OperationFactory(code_operation='ACT1', ventilation_mode='none')
        op.metriques.add(MetriqueFactory(id_indicateur=IndicateurFactory(id_ne=ne)))
        OperationAnneeFactory(id_operation=op, annee=2027, periodicite=True,
                              budget=Decimal('900.00'))
        return plan

    def test_fiche_action_affiche_les_annees_de_prolongation(self, plan):
        ws = load_workbook(io.BytesIO(build_fiche_action_workbook(plan)))['ACT1']
        entete = next(r for r in range(1, ws.max_row + 1)
                      if ws.cell(r, 1).value == 'Programmation annuelle')
        annees = [ws.cell(entete, c).value for c in range(4, ws.max_column + 1)]

        assert annees == list(range(2021, 2028))
        # Case « Périodicité » cochée sous 2027 (colonne D + 6).
        assert ws.cell(entete + 1, 4 + 6).value == 'x'

    def test_budget_compte_les_annees_de_prolongation(self, plan):
        pf = build_plan_finance(plan)
        assert pf.years[-1] == 2027
        assert pf.actions[0].year_total(2027).tot == Decimal('900.00')

    def test_bilan_series_couvre_la_prolongation(self, plan):
        client = APIClient()
        client.force_authenticate(user=SuperAdminFactory())
        resp = client.get(f'/api/plans/realisations/bilan-series/{plan.id_pg}/')
        assert resp.status_code == 200
        assert resp.json()['years'] == list(range(2021, 2028))

    def test_bilan_annonce_l_echeance_effective(self, plan):
        client = APIClient()
        client.force_authenticate(user=SuperAdminFactory())
        resp = client.get(f'/api/plans/realisations/bilan/{plan.id_pg}/')
        assert resp.status_code == 200
        assert resp.json()['annee_max'] == 2027

    def test_api_plan_expose_l_echeance_effective(self, plan):
        client = APIClient()
        client.force_authenticate(user=SuperAdminFactory())
        resp = client.get(f'/api/plans/plans/by-slug/{plan.slug}/')
        assert resp.status_code == 200
        assert resp.json()['annee_fin'] == 2025
        assert resp.json()['annee_fin_effective'] == 2027
