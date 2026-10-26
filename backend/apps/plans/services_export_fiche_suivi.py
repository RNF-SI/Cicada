"""
Export Excel « Fiche de suivi » d'une action (#727) — prévu / réalisé d'UNE année.

Une fiche (onglet) par action. Depuis les exports du plan, on choisit l'année :
seules les actions dont le suivi est saisi cette année-là sont exportées. Depuis
la saisie du suivi, la fiche de l'action pour l'année ouverte (même vide).

Chaque fiche met en regard, colonne « Prévu » / colonne « Réalisé » :
identification (enjeu, indicateur d'état ou de pression, OLT / OO), niveau de
réalisation, opérateurs et financeurs, temps de travail par poste, budget par
poste de dépense, indicateurs de réponse (valeur cible / valeur saisie),
résultat de l'indicateur d'état ou de pression (grille des scores), emprise
prévue et réalisée sur une même carte, commentaires.

Réutilise le rédacteur, le cadre et la carte de la fiche action
(`services_export_fiche_action`) et les postes de dépense de l'export budget
suivi (`services_export_budget_rh`) : mêmes calculs, mêmes libellés.

Point d'entrée public : :func:`build_fiche_suivi_workbook`.
"""

from __future__ import annotations

import io
from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import services_export_fiche_action as fa

_ZERO = Decimal(0)
# A:C libellé · D:F prévu · G:I réalisé. Six colonnes de données : une par
# palier de la grille des scores (indéterminé + 5 niveaux, #727).
_NCOLS = 9
_PREVU, _REALISE = (4, 6), (7, 9)
_SANS_OBJET = "/"    # prévu annuel d'une action programmée globalement (#749)
_SCORE_COLS = [(0, "Indéterminé", "FFDADADA")] + [
    (lvl, fa._SCORE_LABELS[lvl], fa._SCORE_FILLS[lvl]) for lvl in range(1, 6)]


# ---------------------------------------------------------------------------
# Données
# ---------------------------------------------------------------------------

def _annee(op, annee):
    from .models_operations import OperationAnnee
    return (OperationAnnee.objects.filter(id_operation=op, annee=annee)
            .prefetch_related("organismes__realisation").first())


def _realisation(oa):
    try:
        return oa.realisation if oa is not None else None
    except Exception:   # RelatedObjectDoesNotExist
        return None


def _a_un_suivi(oa) -> bool:
    if oa is None:
        return False
    return _realisation(oa) is not None or any(
        getattr(o, "realisation", None) is not None for o in oa.organismes.all())


def actions_suivies(plan, annee, operation_ids=None):
    """Actions à exporter : celles dont le suivi est saisi en `annee`, ou les
    actions demandées (`operation_ids`, export depuis la saisie du suivi)."""
    ops, codes = fa.actions_a_exporter(plan, operation_ids)
    if operation_ids is None:
        ops = [op for op in ops if _a_un_suivi(_annee(op, annee))]
    return ops, codes


def _rh_lignes(op, oa, real):
    """{cible: (jours prévus, jours réalisés)}, dans l'ordre de rencontre."""
    from .services_export_budget_rh import _rh_cible

    out = defaultdict(lambda: [_ZERO, _ZERO])

    def cle(ligne):
        label = _rh_cible(ligne)
        return label if ligne.finance else f"{label} (bénévolat / partenariat)"

    if oa is not None:
        for l in oa.rh_lignes.select_related("id_poste", "id_organisme").order_by("pk"):
            out[cle(l)][0] += l.jours or _ZERO
    if real is not None:
        for l in real.rh_lignes.select_related("id_poste", "id_organisme").order_by("pk"):
            out[cle(l)][1] += l.jours or _ZERO
    return out


def _derniere_mesure(met, annee):
    from .models_indicateurs import Mesure
    return (Mesure.objects.filter(id_metrique=met, date_mesure__year=annee)
            .order_by("-date_mesure", "-date_ajout").first())


def _metriques_du_cadre(op):
    """Métriques des indicateurs d'état / de pression de l'action (#727) :
    celles que l'action vise, sinon toutes celles de l'indicateur."""
    visees = {m.pk for m in op.metriques.all()}
    out = []
    for ind in fa._linked_indicateurs(op):
        if fa._is_reponse(ind):
            continue
        mets = list(ind.metriques.all())
        mets = [m for m in mets if m.pk in visees] or mets
        out.extend((ind, m) for m in mets)
    return out


def _score(mesure, met):
    """Palier 0 (indéterminé) à 5 d'une mesure, calculé comme à l'écran."""
    # serializers_enjeux d'abord : importer views_indicateurs à froid bute
    # sur le cycle serializers_indicateurs ↔ serializers_enjeux.
    from . import serializers_enjeux  # noqa: F401
    from .views_indicateurs import _mesure_to_score
    try:
        return _mesure_to_score(mesure, met) or 0
    except Exception:
        return 0


# ---------------------------------------------------------------------------
# Rendu
# ---------------------------------------------------------------------------

def _cmp(w, label, prevu, realise, *, total=False):
    """Ligne « libellé (A:C) | prévu (D:F) | réalisé (G:I) »."""
    ws = w.ws
    fill = fa._TOTAL_FILL if total else fa._SUBLABEL_FILL
    ws.merge_cells(start_row=w.r, start_column=1, end_row=w.r, end_column=3)
    lc = ws.cell(w.r, 1, label)
    lc.fill = PatternFill("solid", fgColor=fill)
    lc.font = fa._F_TOTAL if total else fa._F_SUBLABEL
    lc.alignment = fa._AL_L
    for (col, fin), v in ((_PREVU, prevu), (_REALISE, realise)):
        ws.merge_cells(start_row=w.r, start_column=col, end_row=w.r, end_column=fin)
        c = ws.cell(w.r, col, v if v not in (None, "") else "—")
        c.font = fa._F_TOTAL if total else fa._F_VALUE
        c.alignment = fa._AL_R if isinstance(v, str) and v[:1].isdigit() else fa._AL_L
        if total:
            c.fill = PatternFill("solid", fgColor=fa._TOTAL_FILL)
    for col in range(1, _NCOLS + 1):
        ws.cell(w.r, col).border = fa._B
    w.r += 1


def _cmp_header(w, label="", prevu="Prévu", realise="Réalisé"):
    ws = w.ws
    ws.merge_cells(start_row=w.r, start_column=1, end_row=w.r, end_column=3)
    lc = ws.cell(w.r, 1, label)
    lc.fill = PatternFill("solid", fgColor=fa._LABEL_FILL)
    lc.font = fa._F_LABEL
    lc.alignment = fa._AL_L
    for (col, fin), text in ((_PREVU, prevu), (_REALISE, realise)):
        ws.merge_cells(start_row=w.r, start_column=col, end_row=w.r, end_column=fin)
        c = ws.cell(w.r, col, text)
        c.fill = PatternFill("solid", fgColor=fa._YEARHDR_FILL)
        c.font = fa._F_YEAR
        c.alignment = fa._AL_C
    for col in range(1, _NCOLS + 1):
        ws.cell(w.r, col).border = fa._B
    w.r += 1


def _grille_scores(w, titre, lignes):
    """Grille « métrique × palier » (#727) : la mesure de l'année est rangée
    dans la colonne de son score. `lignes` : [(libellé, palier, texte)]."""
    ws = w.ws
    ws.merge_cells(start_row=w.r, start_column=1, end_row=w.r, end_column=3)
    lc = ws.cell(w.r, 1, titre)
    lc.fill = PatternFill("solid", fgColor=fa._LABEL_FILL)
    lc.font = fa._F_LABEL
    lc.alignment = fa._AL_L
    for i, (_lvl, label, fill) in enumerate(_SCORE_COLS):
        c = ws.cell(w.r, 4 + i, label)
        c.fill = PatternFill("solid", fgColor=fill)
        c.font = fa._F_SCOREHDR
        c.alignment = fa._AL_C
    for col in range(1, _NCOLS + 1):
        ws.cell(w.r, col).border = fa._B
    ws.row_dimensions[w.r].height = 22
    w.r += 1
    for label, palier, texte in lignes:
        ws.merge_cells(start_row=w.r, start_column=1, end_row=w.r, end_column=3)
        lc = ws.cell(w.r, 1, label)
        lc.fill = PatternFill("solid", fgColor=fa._SUBLABEL_FILL)
        lc.font = fa._F_SUBLABEL
        lc.alignment = fa._AL_L
        if palier is None:
            ws.merge_cells(start_row=w.r, start_column=4, end_row=w.r, end_column=_NCOLS)
            c = ws.cell(w.r, 4, texte)
            c.font = fa._F_VALUE
            c.alignment = fa._AL_C
        else:
            c = ws.cell(w.r, 4 + palier, texte)
            c.font = fa._F_VALUE
            c.alignment = fa._AL_C
            c.fill = PatternFill("solid", fgColor=_SCORE_COLS[palier][2])
        for col in range(1, _NCOLS + 1):
            ws.cell(w.r, col).border = fa._B
        w.r += 1


def _titre(w, texte, sous_titre):
    ws = w.ws
    for text, font, fill in ((texte, fa._F_TITLE, fa._PRIMARY), (sous_titre, fa._F_LABEL, None)):
        ws.merge_cells(start_row=w.r, start_column=1, end_row=w.r, end_column=_NCOLS)
        c = ws.cell(w.r, 1, text)
        c.font = font
        c.alignment = fa._AL_L
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
            ws.row_dimensions[w.r].height = 24
        w.r += 1
    w.blank()


def _render(ws, plan, op, annee, code):
    from .services_export_budget_rh import _budget_postes
    from .services_export_finance import build_action_finance, poste_entry_factory

    for col in range(1, _NCOLS + 1):
        ws.column_dimensions[get_column_letter(col)].width = 14 if col <= 3 else 12
    # L'impression tient en largeur de page (comme la fiche action, #721).
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    w = fa._Writer(ws, _NCOLS)
    oa = _annee(op, annee)
    real = _realisation(oa)
    globale = bool(getattr(op, "programmation_globale", False))
    cadre = fa._cadre(op)
    l_ind, _l_exig, l_obj = fa.cadre_labels(cadre, fa._is_cs(op))

    _titre(w, f"FICHE DE SUIVI — {code} {fa._txt(op.libelle)} — {annee}".replace("  ", " "),
           f"Plan de gestion : {fa._txt(plan.nom)}")

    # --- Identification ---
    w.section("Identification de l'action")
    w.kv("Code action", code or "—")
    w.kv("Enjeu / FCR", " ; ".join(cadre["enjeux"]) or "—")
    # #727 (retour 23/10) — rappel de l'indicateur d'état ou de pression, et
    # ligne d'objectif nommée selon la branche (#766).
    w.kv(l_ind, " ; ".join(cadre["indicateurs"]) or "—")
    w.kv(l_obj, " ; ".join(cadre["objectifs"]) or "—")
    priorite = getattr(op, "id_priorite", None)
    w.kv("Priorité", fa._txt(getattr(priorite, "label", "")) or "—")
    if globale:
        prog = "Programmée globalement sur la période de l'action"
    elif oa is not None and oa.periodicite:
        prog = f"Programmée en {annee}"
    else:
        prog = f"Non programmée en {annee} (réalisée non prévue)"
    w.kv("Programmation", prog)

    # --- Réalisation ---
    w.section(f"Réalisation {annee}")
    niveau = getattr(real, "id_niveau_realisation", None) if real else None
    w.kv("Niveau de réalisation", fa._txt(getattr(niveau, "label", "")) or "Non saisi")
    _cmp_header(w)
    _cmp(w, "Opérateurs", fa._txt(op.operateurs), fa._txt(getattr(real, "operateurs_realises", "")))
    _cmp(w, "Financeurs", fa._txt(op.financeurs), fa._txt(getattr(real, "financeurs_realises", "")))

    # --- Temps de travail ---
    w.section("Temps de travail (jours)")
    _cmp_header(w, "Poste / temps de travail")
    lignes = _rh_lignes(op, oa, real)
    tp = tr = _ZERO
    for label, (p, r) in lignes.items():
        _cmp(w, label, _SANS_OBJET if globale else fa._jours(p), fa._jours(r))
        tp += p; tr += r
    if not lignes:
        _cmp(w, "Aucun temps de travail", "—", "—")
    _cmp(w, "Total temps de travail", _SANS_OBJET if globale else fa._jours(tp),
         fa._jours(tr), total=True)

    # --- Budget ---
    w.section("Budget (€)")
    _cmp_header(w, "Poste de dépense")
    af = build_action_finance(op, {}, defaultdict(poste_entry_factory))
    cell = af.year_total(annee)
    for label, prev_of, real_of in _budget_postes(af):
        p, r = prev_of(cell), real_of(cell)
        if p or r:
            _cmp(w, label, _SANS_OBJET if globale else fa._euro(p), fa._euro(r))
    _cmp(w, "Total budget", _SANS_OBJET if globale else fa._euro(cell.tot),
         fa._euro(cell.rtot), total=True)

    # --- Indicateurs de réponse ---
    w.section("Indicateurs de réponse")
    metriques = [(ind, met) for ind in fa._reponse_indicateurs(op) for met in ind.metriques.all()]
    if metriques:
        _cmp_header(w, "Métrique (indicateur)", "Valeur cible", f"Valeur {annee}")
        for ind, met in metriques:
            unite = f" {fa._txt(met.unite)}" if fa._txt(getattr(met, "unite", "")) else ""
            mesure = _derniere_mesure(met, annee)
            cible = fa._txt(met.etat_reference)
            _cmp(w, f"{fa._txt(met.nom_metrique)} ({fa._txt(ind.nom_indicateur)})",
                 f"{cible}{unite}" if cible else "—",
                 f"{fa._txt(mesure.valeur)}{unite}" if mesure and fa._txt(mesure.valeur) else "—")
    else:
        w.kv("Indicateurs de réponse", "Aucun indicateur de réponse rattaché à l'action")

    # --- Résultat de l'indicateur d'état / de pression (#727) ---
    w.section(f"Résultat : {l_ind[0].lower()}{l_ind[1:]} en {annee}")
    lignes = []
    for ind, met in _metriques_du_cadre(op):
        label = f"{fa._txt(met.nom_metrique)} ({fa._txt(ind.nom_indicateur)})"
        mesure = _derniere_mesure(met, annee)
        if mesure is None or not fa._txt(mesure.valeur):
            lignes.append((label, None, f"Aucune mesure en {annee}"))
            continue
        unite = f" {fa._txt(met.unite)}" if fa._txt(getattr(met, "unite", "")) else ""
        lignes.append((label, _score(mesure, met), f"{annee} : {fa._txt(mesure.valeur)}{unite}"))
    if lignes:
        _grille_scores(w, f"Scores des indicateurs ({l_ind.replace('Indicateur ', '')})", lignes)
    else:
        w.kv(l_ind, "Aucun indicateur d'état ou de pression rattaché à l'action")

    # --- Emprise ---
    w.section("Emprise")
    prevue = fa._operation_geom(op)
    realisee = getattr(real, "geom_realisee", None) if real else None
    png = fa._geom_png(prevue, realise=realisee) if (prevue or realisee) else None
    if png:
        w.picture_row("Emprise prévue (bleu-vert)\net réalisée (terra cotta)", png)
    else:
        w.kv("Emprise", "Aucune emprise prévue ni réalisée")

    # --- Commentaires ---
    w.section("Commentaires du suivi")
    w.kv("Commentaires", fa._txt(getattr(real, "commentaires", "")) or "—")
    ws.sheet_view.showGridLines = False


def build_fiche_suivi_workbook(plan, annee, operation_ids=None) -> bytes:
    """Classeur des fiches de suivi de l'année `annee` (un onglet par action)."""
    fa._appliquer_couleur_instance()
    annee = int(annee)
    ops, codes = actions_suivies(plan, annee, operation_ids)

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    if not ops:
        ws = wb.create_sheet(fa._sanitize_title(f"Suivi {annee}", used))
        ws["A1"] = f"Aucun suivi saisi en {annee} pour ce plan."
        ws["A1"].font = Font(bold=True, size=12, color=fa._PRIMARY)
    for op in ops:
        code = codes.get(op.id_operation) or fa._txt(op.code_operation)
        # Code court (CS1…) d'abord : l'année doit rester lisible dans les
        # 31 caractères d'un nom d'onglet.
        titre = code or fa._txt(op.code_operation) or f"Action {op.id_operation}"
        ws = wb.create_sheet(fa._sanitize_title(f"{titre} {annee}"[-31:] if len(titre) > 25
                                                else f"{titre} {annee}", used))
        _render(ws, plan, op, annee, code)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
